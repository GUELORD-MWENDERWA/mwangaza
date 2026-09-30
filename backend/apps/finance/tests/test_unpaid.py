"""Impayés (G1) : isolation entre écoles, droit `finance.view_reports`, retard après réductions et
sans les paiements annulés, par devise jamais additionnée, filtres, exports tracés."""

import io
from decimal import Decimal

import openpyxl
import pytest

from apps.audit.actions import Action
from apps.audit.models import AuditLog
from apps.core.tenancy import use_school
from apps.finance import services
from apps.finance.tests.test_balance import at
from apps.finance.tests.test_payments import pay, school_with_fees
from apps.students.tests.conftest import enrolled

pytestmark = pytest.mark.django_db

UNPAID = "/api/unpaid/"


@pytest.fixture
def alpha(school, head):
    return school_with_fees(school, head)


@pytest.fixture
def late(monkeypatch, school, head, cashier, alpha):
    """Le 16/10 (minerval 3 x 120,00 $ dès le 15/10, fonctionnement 25 000 FC au 30/09) :
    Kasereka n'a rien payé ; Amani a une bourse de 25 % sur le minerval, a payé 10 000 FC, puis
    90,00 $ annulés ; Paluku a tout payé."""
    group = alpha["group"]
    students = [enrolled(school, head, group, last_name=name) for name in ("Amani", "Paluku")]
    with use_school(school):
        amani, paluku = (student.enrollments.get() for student in students)
        services.grant_discount(
            amani,
            actor=head,
            fee_type=alpha["fee_type"],
            kind="percent",
            percent=25,
            reason_code="scholarship",
            reason="Bourse de la paroisse",
        )
    pay(school, cashier, amani, 10000, "CDF")
    cancelled = pay(school, cashier, amani, 9000, "USD")
    pay(school, cashier, paluku, 12000, "USD")
    pay(school, cashier, paluku, 25000, "CDF")
    with use_school(school):
        request, _created = services.request_cancellation(
            cancelled, actor=cashier, reason="Mauvais élève sélectionné (homonyme)"
        )
        services.approve_cancellation(request, actor=head)
    at(monkeypatch, "2026-10-16")
    return {"kasereka": alpha["enrollment"], "amani": amani, "paluku": paluku}


def amounts(row: dict) -> list:
    return [(amount["currency"], amount["overdue_minor"]) for amount in row["amounts"]]


def test_late_is_after_discounts_without_cancelled_payments_and_per_currency(
    monkeypatch, api, head, late
):
    client = api(head)
    rows = client.get(UNPAID).json()["results"]

    assert [(row["enrollment"], amounts(row), row["days_late"]) for row in rows] == [
        (str(late["kasereka"].pk), [("USD", 12000), ("CDF", 25000)], 16),
        (str(late["amani"].pk), [("USD", 9000), ("CDF", 15000)], 16),
    ]
    # Dernier paiement : pas celui qui est annulé.
    assert rows[1]["last_payment"]["received_currency"] == "CDF"
    assert rows[1]["guardian"]["phone"]

    summary = client.get(f"{UNPAID}summary/").json()
    assert (summary["students"], summary["enrolled"], summary["long_overdue"]) == (2, 3, 0)
    assert amounts(summary) == [("USD", 21000), ("CDF", 40000)]
    assert summary["installments"] == ["Fonctionnement", "1re tranche", "2e tranche", "3e tranche"]

    # Filtres : devise, tranche, retard de plus de 7 jours (seul le fonctionnement, 16 jours).
    in_francs = client.get(UNPAID, currency="CDF", order="amount").json()["results"]
    assert [amounts(row) for row in in_francs] == [[("CDF", 25000)], [("CDF", 15000)]]
    first = client.get(UNPAID, installment="1re tranche").json()["results"]
    assert [amounts(row) for row in first] == [[("USD", 12000)], [("USD", 9000)]]
    week = client.get(f"{UNPAID}summary/", over_days=7).json()
    assert amounts(week) == [("CDF", 40000)]

    at(monkeypatch, "2026-11-01")
    assert client.get(f"{UNPAID}summary/").json()["long_overdue"] == 2


def test_excel_has_one_line_per_late_fee_and_each_export_is_audited(api, school, head, late):
    client = api(head)
    response = client.get(f"{UNPAID}export/", output="xlsx", order="name")
    sheet = openpyxl.load_workbook(io.BytesIO(response.content)).active
    lines = [
        (row[2].value, row[6].value, row[10].value, row[10].number_format, row[11].value)
        for row in sheet.iter_rows(min_row=2)
    ]
    assert lines == [
        ("AMANI Jean", "Fonctionnement", 15000, "#,##0", "FC"),
        ("AMANI Jean", "1re tranche", Decimal("90.00"), "#,##0.00", "$"),
        ("KASEREKA Jean", "Fonctionnement", 25000, "#,##0", "FC"),
        ("KASEREKA Jean", "1re tranche", Decimal("120.00"), "#,##0.00", "$"),
    ]
    preview = client.get(f"{UNPAID}export/", output="html", detail="true")
    assert preview["Content-Type"].startswith("text/html")
    assert "AMANI Jean" in preview.content.decode()
    assert "PALUKU" not in preview.content.decode()

    pdf = client.get(f"{UNPAID}export/", output="pdf", group="list")
    assert pdf["Content-Type"] == "application/pdf"

    with use_school(school):
        logs = AuditLog.objects.filter(action=Action.FINANCE_UNPAID_EXPORTED).order_by("id")
        # L'aperçu à l'écran n'est pas noté ; les fichiers Excel et PDF le sont.
        assert [(log.actor, log.after["output"], log.after["students"]) for log in logs] == [
            (head, "xlsx", 2),
            (head, "pdf", 2),
        ]


def test_a_school_neither_lists_nor_exports_the_unpaid_of_another(
    monkeypatch, api, school, other_school, head, other_head, alpha
):
    beta = school_with_fees(other_school, other_head)
    at(monkeypatch, "2026-10-16")

    rows = api(head).get(UNPAID).json()["results"]
    assert [row["enrollment"] for row in rows] == [str(alpha["enrollment"].pk)]
    other = api(other_head, host="beta.localhost").get(UNPAID).json()["results"]
    assert [row["enrollment"] for row in other] == [str(beta["enrollment"].pk)]

    sheet = openpyxl.load_workbook(
        io.BytesIO(api(head).get(f"{UNPAID}export/", output="xlsx").content)
    ).active
    assert sheet.max_row == 3  # titres + minerval et fonctionnement de son seul élève
    # La classe d'une autre école n'est pas un filtre possible.
    response = api(head).get(UNPAID, class_group=str(beta["group"].pk))
    assert response.status_code == 400


def test_unpaid_need_the_reports_right(api, cashier, alpha):
    client = api(cashier)
    for path in (UNPAID, f"{UNPAID}summary/", f"{UNPAID}export/?output=xlsx"):
        assert client.get(path).status_code == 403

"""Documents à imprimer ou à télécharger (D5, M6) : gabarit HTML rendu tel quel (aperçu) ou en PDF
(WeasyPrint), classeurs Excel (openpyxl). Les images (logo, photo) sont intégrées au HTML en
`data:` : WeasyPrint ne charge rien d'autre, ni fichier ni adresse (un SVG piégé ne peut pas faire
appeler une adresse par le serveur)."""

import base64
import io
import re
from dataclasses import dataclass
from functools import cache
from pathlib import Path

import openpyxl
import segno
from django.http import HttpResponse
from django.template.loader import render_to_string
from django.utils import timezone
from drf_spectacular.types import OpenApiTypes
from drf_spectacular.utils import OpenApiResponse
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from rest_framework import serializers

from apps.core import kept, printing
from apps.core.models import PlatformSettings
from apps.schools import signatures
from apps.schools.logos import CONTENT_TYPES as LOGO_TYPES
from apps.schools.logos import logo_kind
from apps.schools.models import School

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PDF = "application/pdf"
HTML = "text/html"
XLSX_RESPONSE = {(200, XLSX): OpenApiResponse(OpenApiTypes.BINARY, description="Classeur Excel.")}
PRINT_RESPONSE = {
    (200, PDF): OpenApiResponse(OpenApiTypes.BINARY, description="Document PDF (`output=pdf`)."),
    (200, HTML): OpenApiResponse(OpenApiTypes.STR, description="Aperçu HTML (`output=html`)."),
}


PAPERS = ("a4", "a5")
ORIENTATIONS = ("portrait", "landscape")
MAX_PER_SHEET = 3
# Corps réduit quand plusieurs exemplaires se partagent une feuille.
PER_SHEET_FONT = {2: "9pt", 3: "8pt"}


class PrintOptionsSerializer(serializers.Serializer):
    """Choix de la personne qui imprime (8.3.4) : papier et sens ; sans sens, celui du document."""

    paper = serializers.ChoiceField(choices=PAPERS, default="a4")
    orientation = serializers.ChoiceField(
        choices=ORIENTATIONS, required=False, help_text="Absent : le sens habituel du document."
    )


class SlipPrintOptionsSerializer(PrintOptionsSerializer):
    """Documents courts (convocation, circulaire) : plusieurs exemplaires par feuille."""

    per_sheet = serializers.IntegerField(
        default=1, min_value=1, max_value=MAX_PER_SHEET, help_text="Exemplaires par feuille."
    )


def print_options(data: dict, orientation: str = "portrait") -> dict:
    """Options validées → contexte des gabarits (`page_size` pour `@page`, `per_sheet`)."""
    paper = data.get("paper", "a4")
    chosen = data.get("orientation") or orientation
    return {
        "paper": paper,
        "orientation": chosen,
        "page_size": f"{paper.upper()} {chosen}",
        "per_sheet": data.get("per_sheet", 1),
    }


def _repeat_body(html: str, count: int) -> str:
    """Le corps du document `count` fois sur la feuille, séparé d'un filet pointillé."""
    start = html.index("<body>") + len("<body>")
    end = html.rindex("</body>")
    copy = f'<div class="copy">{html[start:end]}</div>'
    style = f"<style>body {{ font-size: {PER_SHEET_FONT[count]}; }}</style>"
    return f"{html[:start]}{style}{copy * count}{html[end:]}"


FONTS = Path(__file__).parent / "static" / "core" / "site"


@cache
def print_fonts() -> str:
    """Police de la marque (Atkinson Hyperlegible Next) intégrée aux documents en `data:` : le PDF
    a la même typographie que l'application, sans que WeasyPrint charge un fichier."""
    rules = []
    for weight in (400, 700, 800):
        path = FONTS / f"atkinson-hyperlegible-next-latin-{weight}-normal.woff2"
        data = base64.b64encode(path.read_bytes()).decode()
        rules.append(
            '@font-face { font-family: "Atkinson Hyperlegible Next"; '
            f'font-weight: {weight}; src: url(data:font/woff2;base64,{data}) format("woff2"); }}'
        )
    return "\n".join(rules)


SERIF_FONTS = Path(__file__).parent / "static" / "core" / "fonts"


@cache
def serif_fonts() -> str:
    """Police à empattements (Liberation Serif, métrique de Times) du bulletin : le formulaire
    officiel est composé en Times, le PDF s'en approche sans dépendre des polices du serveur."""
    rules = []
    for weight, name in ((400, "Regular"), (700, "Bold")):
        data = base64.b64encode((SERIF_FONTS / f"LiberationSerif-{name}.ttf").read_bytes()).decode()
        rules.append(
            '@font-face { font-family: "Liberation Serif"; '
            f'font-weight: {weight}; src: url(data:font/ttf;base64,{data}) format("truetype"); }}'
        )
    return "\n".join(rules)


def data_uri(field, kind: str | None, content_types: dict[str, str]) -> str:
    """Fichier (logo, photo) intégré au document ; vide s'il manque ou n'est plus lisible."""
    if not field or kind is None:
        return ""
    try:
        with field.open("rb") as file:
            content = file.read()
    except OSError:
        return ""
    return f"data:{content_types[kind]};base64,{base64.b64encode(content).decode()}"


def qr_code(text: str) -> str:
    """QR code d'une adresse en SVG `data:`, net à toute taille dans le PDF (`segno`)."""
    return segno.make(text, error="m").svg_data_uri(border=0)


def school_header(school: School) -> dict:
    """En-tête commun : nom officiel (sinon nom) et logo de l'école, date d'impression ; pour le
    bas du document, cachet et chef d'établissement qui signe (`signatures`, DOC.1)."""
    kind = None
    if school.logo:
        try:
            with school.logo.open("rb") as file:
                kind = logo_kind(file)
        except OSError:
            kind = None
    return {
        "school_name": school.official_name or school.name,
        "school_code": school.code,
        "header_lines": [
            line for line in (school.receipt or {}).get("header", "").splitlines() if line.strip()
        ],
        "logo": data_uri(school.logo, kind, LOGO_TYPES),
        "printed_at": timezone.localtime(),
        **signatures.signature_context(school),
    }


def render(
    template: str,
    context: dict,
    output: str,
    name: str,
    options: dict | None = None,
    more: list[dict] | None = None,
    keep: kept.Keep | None = None,
) -> HttpResponse:
    """`html` : aperçu dans la page (même origine seulement) ; `pdf` : fichier ouvert dans le
    navigateur, qui l'imprime ou l'enregistre. `options` : `print_options(...)`. `more` : autres
    contextes du même gabarit (bulletins de toute la classe), chacun sur sa propre page. `keep` :
    chaque page est un PDF gardé (PERF.2), refait seulement si ses données ont changé."""
    options = options or print_options({})
    # Les polices ne vont dans le HTML que pour l'aperçu du navigateur ; le PDF les reçoit à part.
    preview = output == "html"
    shared = {
        "platform": PlatformSettings.load(),
        "print_fonts": print_fonts() if preview else "",
        "serif_fonts": serif_fonts() if preview else "",
        **options,
    }
    if keep and output == "pdf" and options["per_sheet"] == 1:
        return _render_kept(template, [context, *(more or [])], shared, name, keep)
    html = render_to_string(template, {**shared, **context})
    bodies: list[str] = []
    if more:
        for extra in more:
            page = render_to_string(template, {**shared, **extra})
            bodies.append(page[page.index("<body>") + 6 : page.rindex("</body>")])
    if output == "pdf" and options["per_sheet"] == 1:
        start = html.index("<body>") + 6
        pages = [html[start : html.rindex("</body>")], *bodies]
        # Document lourd : jamais dans la requête web, un travailleur à part le fait (PERF.1).
        if printing.is_heavy(len(pages), sum(map(len, pages))):
            return printing.submit(html[:start], pages, name)
    if bodies:
        end = html.rindex("</body>")
        html = (
            html[:end] + "".join(f'<div class="next">{body}</div>' for body in bodies) + html[end:]
        )
    if options["per_sheet"] > 1:
        html = _repeat_body(html, options["per_sheet"])
    if output == "html":
        response = HttpResponse(html, content_type=f"{HTML}; charset=utf-8")
        response["X-Frame-Options"] = "SAMEORIGIN"
        return response
    return file_response(printing.to_pdf(html), PDF, name, inline=True)


def _render_kept(
    template: str, contexts: list[dict], shared: dict, name: str, keep: kept.Keep
) -> HttpResponse:
    """PDF de pièces gardées une à une : les pages déjà gardées sont lues, seules les autres sont
    mises en page (dans la requête si elles sont peu nombreuses, sinon en tâche de fond)."""
    stable = [
        render_to_string(template, {**shared, **context, **keep.volatile}) for context in contexts
    ]
    keys = [kept.fingerprint(html) for html in stable]
    found = kept.lookup(keys)
    start = stable[0].index("<body>") + 6
    head = stable[0][:start]
    todo = [i for i, key in enumerate(keys) if key not in found]
    entries = [
        {"key": key, "subject": str(subject), "kind": keep.kind, "variant": keep.variant}
        for key, subject in zip(keys, keep.subjects, strict=True)
    ]
    bodies: list[str | None] = [None] * len(keys)
    for i in todo:
        page = render_to_string(template, {**shared, **contexts[i]})
        bodies[i] = page[page.index("<body>") + 6 : page.rindex("</body>")]
    if printing.is_heavy(len(todo), sum(len(bodies[i]) for i in todo)):
        return printing.submit(head, bodies, name, entries)
    pdfs = printing.build_kept(head, bodies, entries, found)
    return file_response(kept.merge(pdfs), PDF, name, inline=True)


def workbook(
    headers: list[str],
    rows,
    *,
    title: str = "",
    widths: list[int] | None = None,
    formats: list[str | None] | None = None,
):
    """Classeur d'une feuille, titres en gras, écrit ligne à ligne (milliers de lignes) ;
    `formats` : format Excel de chaque colonne (`#,##0.00`, `dd/mm/yyyy`…), vide : aucun ; une
    valeur `Formatted` garde le sien."""
    book = openpyxl.Workbook(write_only=True)
    # Excel refuse []:*?/\ dans le nom d'une feuille, et plus de 31 caractères.
    sheet = book.create_sheet(re.sub(r"[\[\]:*?/\\]", " ", title)[:31] or None)
    for index, width in enumerate(widths or [22] * len(headers), start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    header_cells = []
    for header in headers:
        cell = WriteOnlyCell(sheet, value=header)
        cell.font = Font(bold=True)
        header_cells.append(cell)
    sheet.append(header_cells)
    formats = formats or []
    for row in rows:
        row = [
            _formatted(sheet, value, formats[index] if index < len(formats) else None)
            for index, value in enumerate(row)
        ]
        sheet.append(row)
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


@dataclass(frozen=True)
class Formatted:
    """Valeur avec son propre format Excel, quand une colonne en mélange (montants en dollars et
    en francs, chacun dans son format)."""

    value: object
    number_format: str


def _formatted(sheet, value, number_format: str | None):
    if isinstance(value, Formatted):
        value, number_format = value.value, value.number_format
    if number_format is None or value is None:
        return value
    cell = WriteOnlyCell(sheet, value=value)
    cell.number_format = number_format
    return cell


def file_response(content: bytes, content_type: str, name: str, *, inline=False) -> HttpResponse:
    response = HttpResponse(content, content_type=content_type)
    disposition = "inline" if inline else "attachment"
    response["Content-Disposition"] = f'{disposition}; filename="{name}"'
    return response

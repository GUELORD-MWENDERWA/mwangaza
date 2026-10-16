"""Import d'une liste d'élèves (D1) : vérification ligne par ligne, import sans doublon avec
fratries et matricules, isolation entre écoles, droit `students.import`, fichier réaliste.
Annulation (D6) : élèves retirés, élèves gardés (paiement, réduction, échéancier particulier,
compte parent), paiements intacts, journal, appareils prévenus."""

import io
import uuid
from datetime import date, datetime, timedelta
from decimal import Decimal

import openpyxl
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils import timezone

from apps.academics import discipline
from apps.academics.models import Course, Evaluation, Grade
from apps.audit.actions import Action
from apps.audit.models import AuditLog
from apps.core.tenancy import use_school
from apps.finance import services as finance
from apps.finance.models import Payment, PaymentAllocation, StudentCharge
from apps.finance.tests.test_payments import pay
from apps.schools.models import School, Term
from apps.students import import_cancellation, imports, services
from apps.students.demo import import_rows, import_workbook
from apps.students.models import Enrollment, Family, Guardian, Student, StudentImport
from apps.students.tests.conftest import new_student
from apps.sync.models import ChangeLog
from tests.factories import UserFactory

IMPORTS = "/api/student-imports/"
HEADERS = [
    "NOM", "POSTNOM", "PRENOM", "SEXE", "DATE NAISS.", "CLASSE", "NOM DU PARENT", "TEL",
    "MATRICULE", "N° ID", "N° PERM.",
]  # fmt: skip
# Lignes 3 à 11 du fichier (titre en ligne 1, colonnes en ligne 2).
ROWS = [
    ["KAMBALE", "MUHINDO", "Josué", "M", datetime(2010, 5, 3), "3E ELEC A", "Kambale Paluku",
     "0971234567", "", "61022-6100210011533", "2-610021001115"],
    ["KAMBALE", "KAVIRA", "Neema", "F", "21/11/2011", "3ème Électricité A", "KAMBALE PALUKU",
     971234567, ""],
    ["PALUKU", "SIVIHWA", "Daniel", "M", "31/02/2008", "3E ELEC A", "", "", ""],
    ["KAVIRA", "NZIAVAKE", "Esther", "F", "05/06/2011", "4E MEC AUTO B", "", "", ""],
    ["KAMBALE", "MUHINDO", "JOSUÉ", "M", "03/05/2010", "3E ELEC A", "", "", ""],
    ["KASEREKA", "", "Jean", "M", "01/01/2010", "3E ELEC A", "", "", ""],
    ["MUMBERE", "", "Élie", "X", "09/07/2011", "3E ELEC A", "Mumbere Kule", "12345", ""],
    ["BAHATI", "AMANI", "Christian", "M", "", "3E ELEC A", "Bahati Jean", "+243 99 123 45 67",
     "ANCIEN-001"],
    ["SIKULI", "", "Grâce", "F", "17/01/2010", "3E ELEC A", "", "", "ANCIEN-001"],
]  # fmt: skip


def xlsx(rows=ROWS, headers=HEADERS, name="liste.xlsx") -> SimpleUploadedFile:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["ÉCOLE ALPHA — LISTE DES ÉLÈVES (fictive)"])
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    output = io.BytesIO()
    book.save(output)
    return SimpleUploadedFile(name, output.getvalue())


def upload(client, file, **data):
    return client.client.post(IMPORTS, {"file": file, **data}, HTTP_HOST=client.host)


def statuses(client, imp_id) -> dict:
    lines = client.get(f"{IMPORTS}{imp_id}/lines/").json()["results"]
    return {
        line["line"]: (line["status"], sorted(issue["code"] for issue in line["issues"]))
        for line in lines
    }


@pytest.fixture
def known(school, head, alpha):
    """Élève déjà dans l'école : la ligne 8 le répète."""
    return new_student(school, head, birth_date=date(2010, 1, 1))


def test_each_row_is_checked_then_imported_once_with_siblings_and_matricules(
    api, school, head, alpha, known, django_capture_on_commit_callbacks
):
    client = api(head)
    response = upload(client, xlsx())
    assert response.status_code == 201, response.json()
    imp = response.json()
    assert imp["counts"] == {"rows": 9, "ready": 2, "warning": 2, "error": 5}
    assert imp["missing_fields"] == []
    assert imp["mapping"]["birth_date"] == 4 and imp["mapping"]["guardian_phone"] == 7
    assert statuses(client, imp["id"]) == {
        3: ("ready", []),
        4: ("ready", ["same_family"]),
        5: ("error", ["invalid_date"]),
        6: ("error", ["unknown_class"]),
        7: ("warning", ["duplicate_in_file"]),
        8: ("error", ["already_in_school"]),
        9: ("error", ["invalid_phone", "invalid_sex"]),
        10: ("warning", ["missing_birth_date"]),
        11: ("error", ["matricule_repeated"]),
    }

    # La classe inconnue est associée à la main.
    group = alpha["group"]
    response = client.patch(
        f"{IMPORTS}{imp['id']}/", {"class_mapping": {"4E MEC AUTO B": str(group.pk)}}
    )
    assert response.json()["counts"]["error"] == 4
    assert response.json()["importable_count"] == 5

    with django_capture_on_commit_callbacks(execute=True):
        response = client.post(f"{IMPORTS}{imp['id']}/confirm/")
    assert response.status_code == 202 and response.json()["status"] == "importing"
    done = client.get(f"{IMPORTS}{imp['id']}/").json()
    assert done["status"] == "done" and done["created_count"] == 5

    with use_school(school):
        created = Student.objects.exclude(pk=known.pk).select_related("family")
        assert {student.last_name for student in created} == {"KAMBALE", "KAVIRA", "BAHATI"}
        siblings = {student.first_name: student.family for student in created}
        assert siblings["Josué"] == siblings["Neema"] != siblings["Christian"]
        guardian = siblings["Josué"].guardians.get()
        assert (guardian.full_name, guardian.phone, guardian.is_primary) == (
            "Kambale Paluku",
            "+243971234567",
            True,
        )
        josue = next(student for student in created if student.first_name == "Josué")
        assert (josue.national_id, josue.permanent_number) == (
            "61022-6100210011533",
            "2-610021001115",
        )
        numbers = [student.matricule for student in created]
        assert "ANCIEN-001" in numbers and len(set(numbers)) == 5 and all(numbers)
        assert Enrollment.objects.filter(class_group=group).count() == 5
        logs = AuditLog.objects.filter(action=Action.STUDENTS_STUDENT_IMPORTED)
        assert {log.entity_id for log in logs} == {student.pk for student in created}
        assert {log.actor_id for log in logs} == {head.pk}
        done = StudentImport.objects.get(pk=imp["id"])
        # Le fichier est effacé ; restent les lignes non importées, pour le rapport.
        assert done.rows == []
        assert [line["line"] for line in done.report["lines"]] == [5, 8, 9, 11]

    # Le même fichier envoyé de nouveau n'importe personne une seconde fois.
    again = upload(client, xlsx()).json()
    assert again["importable_count"] == 0
    response = client.post(f"{IMPORTS}{again['id']}/confirm/")
    assert response.status_code == 409 and response.json()["code"] == "nothing_to_import"
    response = client.post(f"{IMPORTS}{imp['id']}/confirm/")
    assert response.json()["code"] == "import_closed"


def test_waiting_values_are_no_matricule_and_the_school_may_give_none(
    api, school, head, alpha, django_capture_on_commit_callbacks
):
    """4.13 : « - », « 0 », « néant » dans la colonne ne bloquent plus (c'était « matricule
    répété ») ; une école qui n'attribue pas de matricule n'en reçoit aucun de l'import."""
    School.objects.filter(pk=school.pk).update(auto_matricule=False)
    rows = [
        ["KAMBALE", "", "Josué", "M", "03/05/2010", "3E ELEC A", "", "", "-"],
        ["KAVIRA", "", "Esther", "F", "05/06/2011", "3E ELEC A", "", "", "-"],
        ["PALUKU", "", "Daniel", "M", "01/02/2008", "3E ELEC A", "", "", 0],
        ["SIKULI", "", "Grâce", "F", "17/01/2010", "3E ELEC A", "", "", "Néant"],
        ["BAHATI", "", "Christian", "M", "09/07/2011", "3E ELEC A", "", "", "ANCIEN-7"],
        ["MUMBERE", "", "Élie", "M", "09/08/2011", "3E ELEC A", "", "", "ANCIEN-7"],
    ]
    client = api(head)
    imp = upload(client, xlsx(rows)).json()
    assert statuses(client, imp["id"]) == {
        **{line: ("ready", []) for line in range(3, 8)},
        8: ("error", ["matricule_repeated"]),
    }
    with django_capture_on_commit_callbacks(execute=True):
        client.post(f"{IMPORTS}{imp['id']}/confirm/")
    with use_school(school):
        numbers = sorted(Student.objects.values_list("matricule", flat=True))
    assert numbers == ["", "", "", "", "ANCIEN-7"]


def test_import_of_one_school_is_invisible_to_another(api, school, head, alpha, other_head, beta):
    imp = upload(api(head), xlsx()).json()
    other = api(other_head, host="beta.localhost")
    path = f"{IMPORTS}{imp['id']}/"
    assert other.get(path).status_code == 404
    assert other.get(f"{path}lines/").status_code == 404
    assert other.get(f"{path}report/").status_code == 404
    assert other.patch(path, {"class_mapping": {}}).status_code == 404
    assert other.post(f"{path}confirm/").status_code == 404
    # Ni une classe ni une année d'une autre école.
    response = api(head).patch(path, {"class_mapping": {"4E MEC AUTO B": str(beta["group"].pk)}})
    assert response.status_code == 400 and "class_mapping" in response.json()["fields"]
    response = upload(api(head), xlsx(), academic_year=str(beta["year"].pk))
    assert response.status_code == 400 and "academic_year" in response.json()["fields"]
    with use_school(school):
        assert StudentImport.objects.get().class_mapping == {}


def test_import_needs_its_permission(api, alpha, cashier):
    client = api(cashier)
    assert upload(client, xlsx()).status_code == 403
    assert client.get(f"{IMPORTS}template/").status_code == 403
    assert client.get(IMPORTS).status_code == 403
    path = f"{IMPORTS}{uuid.uuid4()}/"
    assert client.get(f"{path}cancellation/").status_code == 403
    assert client.post(f"{path}cancel/", {"reason": "Mauvaise année"}).status_code == 403


def test_template_and_report_are_read_back_by_the_import(api, head, alpha):
    client = api(head)
    template = client.get(f"{IMPORTS}template/")
    headers = [
        cell.value for cell in openpyxl.load_workbook(io.BytesIO(template.content)).active[1]
    ]
    assert None not in imports.guess_mapping(headers).values()

    imp = upload(client, xlsx()).json()
    report = openpyxl.load_workbook(io.BytesIO(client.get(f"{IMPORTS}{imp['id']}/report/").content))
    rows = list(report.active.iter_rows(values_only=True))
    assert [row[0] for row in rows[1:]] == [5, 6, 7, 9, 10, 11]
    assert rows[1][2:5] == ("PALUKU", "SIVIHWA", "Daniel")
    # Corrigé, le rapport se réimporte tel quel : mêmes colonnes reconnues.
    assert imports.guess_mapping(list(rows[0]))["class_name"] == imp["mapping"]["class_name"] + 2


def test_csv_from_excel_is_read_and_old_excel_refused(api, head, alpha):
    client = api(head)
    content = (
        "NOM;POSTNOM;PRENOM;SEXE;DATE NAISS.;CLASSE\nKAMBALE;MUHINDO;Josué;M;03/05/2010;3E ELEC A\n"
    )
    response = upload(client, SimpleUploadedFile("liste.csv", content.encode("cp1252")))
    assert response.json()["counts"]["ready"] == 1
    response = upload(client, SimpleUploadedFile("liste.xls", b"\xd0\xcf\x11\xe0"))
    assert response.status_code == 400 and response.json()["code"] == "old_excel"


def test_realistic_file_of_800_students_is_imported_in_a_few_queries(
    school, head, alpha, django_assert_max_num_queries
):
    content = import_workbook(import_rows(800, ["3E ELEC A", "3E ELECTRICITE A"]))
    with use_school(school):
        fee_type = finance.create_fee_type(actor=head, name="Minerval", category="tuition")
        finance.create_fee_rule(
            actor=head,
            academic_year=alpha["year"],
            fee_type=fee_type,
            currency="USD",
            total_amount_minor=24000,
            schedule_kind="installments",
            installments=[
                finance.Installment("1re tranche", 12000, date(2026, 10, 15)),
                finance.Installment("2e tranche", 12000, date(2027, 1, 15)),
            ],
        )
        imp = imports.create_import(
            actor=head, file=SimpleUploadedFile("eleves.xlsx", content), academic_year=alpha["year"]
        )
        counts = imp.report["counts"]
        assert counts["rows"] == 800 and counts["error"] >= 3
        StudentImport.objects.filter(pk=imp.pk).update(status=StudentImport.Status.IMPORTING)
        with django_assert_max_num_queries(40):
            imports.run(imp, actor=head)
        imp.refresh_from_db()
        assert imp.created_count == counts["ready"] + counts["warning"] == Student.objects.count()
        assert Family.objects.count() < imp.created_count  # fratries regroupées
        assert Student.objects.values("matricule").distinct().count() == imp.created_count
        assert StudentCharge.objects.count() == 2 * imp.created_count
        # Annulé d'un seul envoi, lui aussi.
        with django_assert_max_num_queries(30):
            import_cancellation.cancel(imp, actor=head, reason="Mauvaise année")
        assert not Student.objects.exists() and not StudentCharge.objects.exists()
        assert not Enrollment.objects.exists() and not Family.objects.exists()


# Annulation d'un import (D6)

# Deux enfants de la même famille que rien ne touche, puis un élève par raison de le garder.
CANCEL_ROWS = [
    ["KAMBALE", "MUHINDO", "Josué", "M", "03/05/2010", "3E ELEC A", "Kambale Paluku", "0971234567"],
    ["KAMBALE", "KAVIRA", "Neema", "F", "21/11/2011", "3E ELEC A", "Kambale Paluku", "0971234567"],
    ["PALUKU", "SIVIHWA", "Daniel", "M", "12/02/2008", "3E ELEC A", "", ""],
    ["KAVIRA", "NZIAVAKE", "Esther", "F", "05/06/2011", "3E ELEC A", "", ""],
    ["BAHATI", "AMANI", "Christian", "M", "17/01/2010", "3E ELEC A", "", ""],
    ["SIKULI", "", "Grâce", "F", "17/03/2010", "3E ELEC A", "Sikuli Jeanne", "0991234567"],
]  # fmt: skip


def minerval(actor, year):
    """Minerval de 240,00 $ en deux tranches pour toute l'école (école courante)."""
    fee_type = finance.create_fee_type(actor=actor, name="Minerval", category="tuition")
    finance.create_fee_rule(
        actor=actor,
        academic_year=year,
        fee_type=fee_type,
        currency="USD",
        total_amount_minor=24000,
        schedule_kind="installments",
        installments=[
            finance.Installment("1re tranche", 12000, date(2026, 10, 15)),
            finance.Installment("2e tranche", 12000, date(2027, 1, 15)),
        ],
    )
    return fee_type


def imported(school, actor, year, rows=CANCEL_ROWS) -> StudentImport:
    """Import fait jusqu'au bout."""
    with use_school(school):
        imp = imports.create_import(actor=actor, file=xlsx(rows, HEADERS[:8]), academic_year=year)
        StudentImport.objects.filter(pk=imp.pk).update(status=StudentImport.Status.IMPORTING)
        imports.run(imp, actor=actor)
        imp.refresh_from_db()
    return imp


def test_cancelling_an_import_removes_untouched_students_and_never_touches_money(
    api, school, head, cashier, alpha
):
    with use_school(school):
        fee_type = minerval(head, alpha["year"])
        finance.set_exchange_rate(actor=head, academic_year=alpha["year"], rate=Decimal("2850"))
        finance.create_cash_register(actor=head, name="Caisse 1", code="C1")
    before = new_student(school, head, last_name="Mumbere", birth_date=date(2009, 4, 2))
    imp = imported(school, head, alpha["year"])
    assert imp.created_count == 6
    with use_school(school):
        enrollment = {e.student.first_name: e for e in Enrollment.objects.select_related("student")}
        finance.grant_discount(
            enrollment["Esther"],
            actor=head,
            fee_type=fee_type,
            kind="percent",
            percent=25,
            reason_code="scholarship",
            reason="Bourse de la paroisse",
        )
        finance.reschedule_charges(
            enrollment["Christian"],
            actor=head,
            fee_type=fee_type,
            installments=[
                finance.Installment("Versement 1", 4000, date(2026, 11, 1)),
                finance.Installment("Versement 2", 20000, date(2027, 2, 1)),
            ],
            reason="Arrangement avec la famille",
            parent_signed=True,
        )
        Family.objects.filter(students__first_name="Grâce").update(
            primary_user=UserFactory(), activated_at=timezone.now()
        )
        removed = {enrollment[name].pk for name in ("Josué", "Neema")}
        gone = {enrollment[name].student_id for name in ("Josué", "Neema")}
    payment = pay(school, cashier, enrollment["Daniel"], 5000, "USD")
    with use_school(school):
        # Une note garde aussi l'élève (8.2.1).
        course = Course.objects.create(class_group=alpha["group"], name="Physique", period_max=20)
        term = Term.objects.filter(academic_year=alpha["year"], kind="period").first()
        evaluation = Evaluation.objects.create(
            course=course, term=term, name="Interro", max_score=10
        )
        Grade.objects.create(evaluation=evaluation, enrollment=enrollment["Daniel"], score=7)
        # Et une faute de conduite (8.5.1a) : sans cela l'annulation de l'import se bloquerait.
        rule = discipline.create_rule(actor=head, label="Retard", points=1)
        discipline.record_incident(
            actor=head,
            enrollment=enrollment["Christian"],
            rule=rule,
            occurred_on=timezone.localdate(),
            term=term,
        )

    client = api(head)
    path = f"{IMPORTS}{imp.pk}/"
    preview = client.get(f"{path}cancellation/").json()
    assert (preview["removable_count"], preview["kept_count"]) == (2, 4)
    assert {student["first_name"]: student["reasons"] for student in preview["kept"]} == {
        "Daniel": ["payment", "grades"],
        "Esther": ["discount"],
        "Christian": ["schedule", "grades"],
        "Grâce": ["portal"],
    }
    response = client.post(f"{path}cancel/", {"reason": " "})
    assert response.status_code == 400 and "reason" in response.json()["fields"]
    with use_school(school):
        assert Student.objects.count() == 7

    # Passé 30 jours, plus d'annulation.
    with use_school(school):
        late = timezone.now() - timedelta(days=31)
        StudentImport.objects.filter(pk=imp.pk).update(finished_at=late)
    assert client.get(path).json()["cancellable_until"] < timezone.now().isoformat()
    response = client.post(f"{path}cancel/", {"reason": "Trop tard"})
    assert response.status_code == 409 and response.json()["code"] == "import_too_old"
    with use_school(school):
        assert Student.objects.count() == 7
        StudentImport.objects.filter(pk=imp.pk).update(finished_at=timezone.now())

    response = client.post(f"{path}cancel/", {"reason": "Fichier importé dans la mauvaise classe"})
    assert response.status_code == 200, response.json()
    done = response.json()
    assert (done["status"], done["cancelled_count"]) == ("cancelled", 2)
    assert done["cancelled_by_name"] == str(head)

    with use_school(school):
        left = {student.first_name: student for student in Student.objects.all()}
        assert set(left) == {"Jean", "Daniel", "Esther", "Christian", "Grâce"}
        assert left["Jean"].pk == before.pk
        assert {s.student_import_id for s in left.values() if s.pk != before.pk} == {imp.pk}
        # La famille des deux enfants retirés part avec son tuteur ; les autres restent.
        assert not Guardian.objects.filter(full_name="Kambale Paluku").exists()
        assert Family.objects.count() == 5
        assert not Enrollment.objects.filter(pk__in=removed).exists()
        assert not StudentCharge.objects.filter(enrollment_id__in=removed).exists()
        # L'argent n'a pas bougé : le paiement, sa répartition, ce que doivent les élèves gardés.
        assert Payment.objects.get().pk == payment.pk
        assert PaymentAllocation.objects.filter(payment=payment).count() == 1
        allocations = PaymentAllocation.objects.filter(payment=payment)
        assert sum(allocations.values_list("amount_minor", flat=True)) == 5000
        assert StudentCharge.objects.filter(enrollment=enrollment["Daniel"]).count() == 2
        # Les appareils retirent les deux inscriptions à leur prochaine synchronisation.
        deleted = ChangeLog.objects.filter(operation=ChangeLog.Operation.DELETE)
        assert {row.entity_id for row in deleted} == removed
        logs = AuditLog.objects.filter(action=Action.STUDENTS_STUDENT_REMOVED)
        assert {log.entity_id for log in logs} == gone
        assert {log.before["reason"] for log in logs} == {"Fichier importé dans la mauvaise classe"}
        log = AuditLog.objects.get(action=Action.STUDENTS_IMPORT_CANCELLED)
        assert (log.actor_id, log.entity_id) == (head.pk, imp.pk)
        assert (log.after["students_removed"], log.after["students_kept"]) == (2, 4)
        # Les numéros de matricule rendus servent de nouveau.
        assert services.next_matricule_number(alpha["year"]) == 1

    # Une seule annulation par import ; le même fichier se ré-importe, sans les élèves gardés.
    response = client.post(f"{path}cancel/", {"reason": "Encore"})
    assert response.status_code == 409 and response.json()["code"] == "import_not_cancellable"
    again = upload(client, xlsx(CANCEL_ROWS, HEADERS[:8])).json()
    assert again["importable_count"] == 2
    listed = client.get(IMPORTS).json()["results"]
    assert [(row["id"], row["status"], row["cancelled_count"]) for row in listed] == [
        (str(imp.pk), "cancelled", 2)
    ]


def test_import_of_one_school_is_not_cancelled_by_another(
    api, school, head, alpha, other_school, other_head, beta
):
    imp = imported(school, head, alpha["year"])
    theirs = imported(other_school, other_head, beta["year"])
    other = api(other_head, host="beta.localhost")
    path = f"{IMPORTS}{imp.pk}/"
    assert other.get(f"{path}cancellation/").status_code == 404
    assert other.post(f"{path}cancel/", {"reason": "Erreur"}).status_code == 404
    assert [row["id"] for row in other.get(IMPORTS).json()["results"]] == [str(theirs.pk)]
    # L'autre école annule le sien : rien ne bouge dans la première.
    response = other.post(f"{IMPORTS}{theirs.pk}/cancel/", {"reason": "Erreur"})
    assert response.json()["cancelled_count"] == 6
    with use_school(other_school):
        assert not Student.objects.exists() and not Family.objects.exists()
    with use_school(school):
        assert Student.objects.count() == 6 and Enrollment.objects.count() == 6
        assert StudentImport.objects.get().status == StudentImport.Status.DONE
        assert not AuditLog.objects.filter(action=Action.STUDENTS_IMPORT_CANCELLED).exists()

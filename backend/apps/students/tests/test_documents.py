"""Impressions (D5, M11) : fiche d'inscription avec le code famille noté dans le journal à chaque
PDF, liste de classe des seuls élèves actifs en PDF ou Excel. Toujours dans l'école de l'adresse."""

import io

import openpyxl
import pytest

from apps.audit.actions import Action
from apps.audit.models import AuditLog
from apps.core.tenancy import use_school
from apps.students.documents import HIDDEN_CODE
from apps.students.models import Enrollment
from apps.students.tests.conftest import enrolled
from tests.factories import MembershipFactory

pytestmark = pytest.mark.django_db

ENROLLMENTS = "/api/enrollments/"
CLASS_LIST = f"{ENROLLMENTS}class-list/"


@pytest.fixture
def teacher(school):
    """Sans `students.view`."""
    return MembershipFactory(school=school, roles=["teacher"]).user


def test_registration_form_prints_the_family_code_only_when_asked_and_logs_each_pdf(
    api, school, head, cashier, teacher, other_head, alpha, beta
):
    student = enrolled(school, head, alpha["group"], birth_place="Goma")
    with use_school(school):
        enrollment = Enrollment.objects.get(student=student)
        code = f"{student.family.code[:4]}-{student.family.code[4:]}"
    form = f"{ENROLLMENTS}{enrollment.pk}/form/"
    client = api(cashier)

    preview = client.get(form, output="html", family_code="true").content.decode()
    third_party = client.get(form, output="html").content.decode()
    pdf = client.get(form, family_code="true", paper="a5")

    # L'aperçu montre la place du code sans le donner ; sans la case, pas de code du tout.
    assert HIDDEN_CODE in preview and code not in preview
    assert student.full_name in preview and "Goma" in preview
    assert HIDDEN_CODE not in third_party and code not in third_party
    assert pdf["Content-Type"] == "application/pdf" and pdf.content.startswith(b"%PDF")
    with use_school(school):
        printed = AuditLog.objects.get(action=Action.STUDENTS_FAMILY_CODE_PRINTED)
    assert (printed.actor, printed.entity_id, printed.entity_label) == (
        cashier,
        student.family_id,
        student.full_name,
    )
    # Sans `students.view` : refusé ; l'autre école ne voit pas l'inscription.
    assert api(teacher).get(form).status_code == 403
    assert api(other_head, host="beta.localhost").get(form).status_code == 404


def test_class_list_has_the_active_students_of_the_class_with_the_chosen_columns(
    api, school, other_school, head, teacher, other_head, alpha, beta
):
    group = alpha["group"]
    enrolled(school, head, group, last_name="Paluku", first_name="Élie", sex="F")
    enrolled(school, head, group, last_name="Kasereka", middle_name="Mumbere")
    gone = enrolled(school, head, group, last_name="Amani", first_name="Prince")
    enrolled(other_school, other_head, beta["group"], last_name="Bahati")
    with use_school(school):
        Enrollment.objects.filter(student=gone).update(status=Enrollment.Status.TRANSFERRED)
    client = api(head)
    query = {"class_group": str(group.pk), "columns": ["guardian_phone"], "blank_columns": 2}

    sheet = openpyxl.load_workbook(
        io.BytesIO(client.get(CLASS_LIST, output="xlsx", **query).content)
    ).active
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    preview = client.get(CLASS_LIST, output="html", **query).content.decode()
    pdf = client.get(CLASS_LIST, class_group=str(group.pk))

    # Parti : absent ; ordre NOM POST-NOM Prénom ; N°, nom, téléphone et deux colonnes vides.
    assert rows == [
        (1, "Kasereka Mumbere Jean", "+243 97 123 45 67", None, None),
        (2, "Paluku Élie", "+243 97 123 45 67", None, None),
    ]
    assert sheet.max_column == 5
    assert "Amani" not in preview and "Bahati" not in preview
    assert pdf["Content-Type"] == "application/pdf" and pdf.content.startswith(b"%PDF")
    assert api(teacher).get(CLASS_LIST, class_group=str(group.pk)).status_code == 403
    # La classe d'une autre école n'existe pas pour celle-ci.
    other = client.get(CLASS_LIST, class_group=str(beta["group"].pk), output="html")
    assert other.status_code == 400 and "Bahati" not in other.content.decode()


def test_class_list_by_student_number(api, school, head, alpha):
    group = alpha["group"]
    first = enrolled(school, head, group, last_name="Zawadi")
    enrolled(school, head, group, last_name="Amani")
    content = (
        api(head)
        .get(CLASS_LIST, class_group=str(group.pk), output="xlsx", order="matricule")
        .content
    )
    names = [
        row[1]
        for row in openpyxl.load_workbook(io.BytesIO(content)).active.iter_rows(
            min_row=2, values_only=True
        )
    ]

    # Zawadi a reçu le premier matricule : il passe devant Amani.
    assert names[0] == first.full_name
